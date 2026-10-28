# core/recibo_generator.py

from __future__ import annotations

from pathlib import Path
from datetime import date
from decimal import Decimal
from typing import List, Optional
import re

try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None
    Pt = None
    WD_ALIGN_PARAGRAPH = None

from .models import (
    LancamentoExtrato,
    DocumentoFinanceiro,
    ResultadoConcilacao,
    StatusConciliacao,
    TipoDocumento,
)


class ReciboGenerator:
    """Gerador de recibos a partir do modelo .docx, preservando formatação."""

    MESES = {
        1: "janeiro",
        2: "fevereiro",
        3: "março",
        4: "abril",
        5: "maio",
        6: "junho",
        7: "julho",
        8: "agosto",
        9: "setembro",
        10: "outubro",
        11: "novembro",
        12: "dezembro",
    }

    TEMPLATE_PREFERIDO = "Dividendo-Recibos.docx"

    def __init__(self, modelo_path: Optional[Path] = None):
        if modelo_path is None:
            modelo_path = self._localizar_modelo()

        self.modelo_path = modelo_path.resolve() if modelo_path else None
        if self.modelo_path and not self.modelo_path.exists():
            self.modelo_path = None

        self.numero_contador = 1

        if self.modelo_path is None:
            self.modelo_path = self._criar_modelo_padrao()

    def _localizar_modelo(self) -> Optional[Path]:
        template_dir = Path(__file__).parent.parent / "templates"
        if not template_dir.exists():
            return None

        preferido = template_dir / self.TEMPLATE_PREFERIDO
        if preferido.exists():
            return preferido

        modelos = sorted(template_dir.glob("*.docx"))
        return modelos[0] if modelos else None

    def _criar_modelo_padrao(self) -> Path:
        """Fallback: cria um modelo simples se o template oficial não existir."""
        if Document is None:
            raise ImportError("python-docx não está instalado.")

        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(11)

        titulo = doc.add_paragraph()
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = titulo.add_run("RECIBO")
        run.bold = True
        run.font.size = Pt(18)

        valor_p = doc.add_paragraph()
        valor_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = valor_p.add_run("R$ {VALOR}")
        run.bold = True
        run.font.size = Pt(18)

        corpo = doc.add_paragraph()
        corpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = corpo.add_run(
            "Recibo de INCA LTDA – Incorporação, Construções e Administração de Imóveis LTDA, "
            "a quantia de R$ {VALOR} ({VALOR_EXTENSO}), referente a {DESCRICAO}"
        )
        run.bold = True
        run.font.size = Pt(16)

        data_p = doc.add_paragraph()
        data_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = data_p.add_run("Manaus, {DATA}.")
        run.bold = True
        run.font.size = Pt(16)

        template_dir = Path(__file__).parent.parent / "templates"
        template_dir.mkdir(exist_ok=True)
        modelo_temp = template_dir / "recibo_modelo.docx"
        doc.save(str(modelo_temp))
        return modelo_temp

    def gerar_recibo(
        self,
        lancamento: LancamentoExtrato,
        comprovante: DocumentoFinanceiro,
        destino: Path,
    ) -> Path:
        """Gera um recibo .docx preenchendo o modelo e preservando a formatação."""
        if Document is None:
            raise ImportError("python-docx não está instalado.")

        favorecido = self._extrair_favorecido(comprovante, lancamento)
        descricao = self._extrair_descricao(comprovante, lancamento, favorecido)
        valor = abs(lancamento.valor)
        data = lancamento.data

        valor_formatado = f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        data_extenso = f"{data.day:02d} de {self.MESES[data.month]} de {data.year}"

        if self.modelo_path and self.modelo_path.exists():
            doc = Document(str(self.modelo_path))
        else:
            doc = Document()

        placeholders = {
            "{VALOR}": valor_formatado,
            "{VALOR_EXTENSO}": self._valor_por_extenso(valor),
            "{DESCRICAO}": descricao,
            "{FAVORECIDO}": favorecido,
            "{DATA}": data_extenso,
            "{CIDADE}": "Manaus",
            "{REFERENCIA}": self._gerar_referencia(lancamento, comprovante),
            "{NUMERO}": f"R{self.numero_contador:06d}",
        }

        self._substituir_placeholders(doc, placeholders)
        self.numero_contador += 1

        destino = destino.with_suffix(".docx")
        destino.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(destino))
        return destino

    def _substituir_placeholders(self, doc, placeholders: dict) -> None:
        for paragraph in doc.paragraphs:
            self._substituir_no_paragraph(paragraph, placeholders)

        for tabela in doc.tables:
            for row in tabela.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._substituir_no_paragraph(paragraph, placeholders)

    def _substituir_no_paragraph(self, paragraph, placeholders: dict) -> None:
        """
        Substitui placeholders preservando bold, tamanho e fonte dos runs.

        O Word costuma fragmentar o texto em vários runs. Juntamos o texto,
        fazemos o replace e recolocamos tudo no primeiro run, copiando o
        estilo de referência — sem usar paragraph.text = ..., que apaga a
        formatação.
        """
        if not paragraph.runs:
            return

        texto_atual = "".join(run.text for run in paragraph.runs)
        if not texto_atual:
            return

        texto_novo = texto_atual
        for placeholder, valor in placeholders.items():
            texto_novo = texto_novo.replace(placeholder, str(valor))

        if texto_novo == texto_atual:
            return

        ref_run = paragraph.runs[0]
        for run in paragraph.runs:
            if run.bold is not None or run.font.size is not None:
                ref_run = run
                break

        bold = ref_run.bold
        size = ref_run.font.size
        name = ref_run.font.name

        paragraph.runs[0].text = texto_novo
        if bold is not None:
            paragraph.runs[0].bold = bold
        if size is not None:
            paragraph.runs[0].font.size = size
        if name is not None:
            paragraph.runs[0].font.name = name

        for run in paragraph.runs[1:]:
            run.text = ""

    def _valor_por_extenso(self, valor: Decimal) -> str:
        """Converte valor monetário para extenso em português (com centavos por extenso)."""
        try:
            valor_str = f"{valor:.2f}"
            inteiro_str, decimal_str = valor_str.split(".")
            inteiro = int(inteiro_str)
            decimal_int = int(decimal_str)

            if inteiro == 0:
                extenso = "Zero"
            else:
                extenso = self._numero_por_extenso(inteiro)

            if decimal_int == 0:
                return f"{extenso} Reais"

            centavos_extenso = self._numero_por_extenso(decimal_int)
            if decimal_int == 1:
                return f"{extenso} Reais e {centavos_extenso} Centavo"
            return f"{extenso} Reais e {centavos_extenso} Centavos"
        except Exception:
            return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    def _numero_por_extenso(self, numero: int) -> str:
        if numero == 0:
            return "Zero"

        unidades = [
            "", "Um", "Dois", "Três", "Quatro", "Cinco",
            "Seis", "Sete", "Oito", "Nove",
        ]
        especiais = [
            "Dez", "Onze", "Doze", "Treze", "Quatorze", "Quinze",
            "Dezesseis", "Dezessete", "Dezoito", "Dezenove",
        ]
        dezenas = [
            "", "", "Vinte", "Trinta", "Quarenta", "Cinquenta",
            "Sessenta", "Setenta", "Oitenta", "Noventa",
        ]
        centenas = [
            "", "Cento", "Duzentos", "Trezentos", "Quatrocentos",
            "Quinhentos", "Seiscentos", "Setecentos", "Oitocentos", "Novecentos",
        ]

        if numero < 10:
            return unidades[numero]
        if 10 <= numero < 20:
            return especiais[numero - 10]
        if numero < 100:
            dezena = numero // 10
            unidade = numero % 10
            if unidade == 0:
                return dezenas[dezena]
            return f"{dezenas[dezena]} e {unidades[unidade]}"
        if numero < 1000:
            centena = numero // 100
            resto = numero % 100
            if resto == 0:
                return "Cem" if centena == 1 else centenas[centena]
            if centena == 1:
                return f"Cento e {self._numero_por_extenso(resto)}"
            return f"{centenas[centena]} e {self._numero_por_extenso(resto)}"
        if numero < 1_000_000:
            milhares = numero // 1000
            resto = numero % 1000
            if milhares == 1:
                if resto == 0:
                    return "Mil"
                return f"Mil e {self._numero_por_extenso(resto)}"
            if resto == 0:
                return f"{self._numero_por_extenso(milhares)} Mil"
            return f"{self._numero_por_extenso(milhares)} Mil e {self._numero_por_extenso(resto)}"
        if numero < 1_000_000_000:
            milhoes = numero // 1_000_000
            resto = numero % 1_000_000
            if milhoes == 1:
                if resto == 0:
                    return "Um Milhão"
                return f"Um Milhão e {self._numero_por_extenso(resto)}"
            if resto == 0:
                return f"{self._numero_por_extenso(milhoes)} Milhões"
            return f"{self._numero_por_extenso(milhoes)} Milhões e {self._numero_por_extenso(resto)}"

        return str(numero)

    def _descricao_util(self, texto: str) -> bool:
        if not texto or len(texto.strip()) < 3:
            return False
        baixo = texto.lower().strip()
        if re.match(r"^e\d+", baixo):
            return False
        # Códigos de identificação Bradesco (E60746948... / longos alfanuméricos sem espaço)
        if re.fullmatch(r"[A-Z0-9]{20,}", texto.strip(), flags=re.IGNORECASE):
            return False
        if re.fullmatch(r"E\d{10,}[A-Z0-9]*", texto.strip(), flags=re.IGNORECASE):
            return False
        # Só valor monetário (campo Descrição vazio no PIX)
        if re.fullmatch(r"R\$\s*[\d.,]+", texto.strip(), flags=re.IGNORECASE):
            return False
        if re.fullmatch(r"[\d.,]+", texto.strip()):
            return False
        ruido = [
            "telefone", "sac", "ouvidoria", "bradesco", "pagamento não identificado",
            "pagamento nao identificado", "pagto efetuado", "comprovante de",
            "bradesco net empresa", "dados da transferência", "identificação",
            "identificacao", "consulte o site", "demais telefones", "fale conosco",
            "cancelamentos", "atendimento", "informações", "informacoes",
            "deficiente auditivo", "alo bradesco", "alô bradesco",
        ]
        if any(p in baixo for p in ruido):
            return False
        if baixo in {"descrição", "descricao", "descrição:", "descricao:"}:
            return False
        if baixo.endswith(":") and len(baixo) < 20:
            return False
        return True

    def _extrair_descricao(
        self,
        comprovante: DocumentoFinanceiro,
        lancamento: LancamentoExtrato,
        favorecido: str,
    ) -> str:
        """
        Prioridade para o recibo:
        1. Campo Descrição no texto do PDF (nunca Identificação)
        2. comprovante.descricao se for texto útil
        3. DIVIDENDO {favorecido}
        """
        texto = comprovante.texto_extraido or ""

        # Somente padrões de DESCRIÇÃO (não usar Identificação — é código do banco)
        candidatos = []
        for padrao in [
            # valor ANTES do rótulo: "PC_237_MATERIAL ITNA03Descrição:"
            r"([A-Za-z0-9][^\n]{2,100}?)\s*Descri[cç][aã]o\s*:",
            # valor DEPOIS do rótulo: "Descrição: ALUGUEL ..."
            r"Descri[cç][aã]o\s*:?\s*([^\n]{3,120})",
        ]:
            for match in re.finditer(padrao, texto, re.IGNORECASE):
                desc = match.group(1).strip()
                desc = re.split(
                    r"(?:Mídia|Valor|Tarifa|Documento|Instituição|Nome|CNPJ|Identifica)\s*:",
                    desc,
                    maxsplit=1,
                    flags=re.IGNORECASE,
                )[0].strip()
                desc = re.sub(r"\s+", " ", desc)
                if self._descricao_util(desc):
                    candidatos.append(desc)

        if candidatos:
            # Prefere o candidato mais "legível" (com espaço ou underscore de PC_)
            candidatos.sort(key=lambda s: (0 if " " in s or "PC_" in s.upper() else 1, -len(s)))
            return candidatos[0][:240]

        if comprovante.descricao and self._descricao_util(comprovante.descricao):
            return comprovante.descricao.strip()[:240]

        nome = (favorecido or "").strip()
        if nome and nome.upper() not in {"NÃO IDENTIFICADO", "NAO IDENTIFICADO"}:
            return f"DIVIDENDO {nome.upper()}"

        return "DIVIDENDO"

    def _extrair_favorecido(
        self,
        comprovante: DocumentoFinanceiro,
        lancamento: LancamentoExtrato,
    ) -> str:
        if comprovante.favorecido and len(comprovante.favorecido.strip()) > 3:
            return comprovante.favorecido.strip()[:120]

        texto = comprovante.texto_extraido or ""
        padroes = [
            r"(?:Dados de quem recebeu\s*Nome:|Nome do favorecido:)\s*([^\n]{5,})",
            r"(?:Nome|FAVORECIDO|Beneficiário)\s*:?\s*([^\n]{5,})",
            r"(?:favorecido|beneficiario|recebedor|fornecedor)\s*:?\s*([^\n]{5,})",
        ]
        for padrao in padroes:
            match = re.search(padrao, texto, re.IGNORECASE)
            if match:
                nome = match.group(1).strip()
                nome = re.sub(r"^Nome:\s*", "", nome, flags=re.IGNORECASE)
                nome = re.sub(r"\s*\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", "", nome)
                if len(nome) > 5 and not re.match(r"^[\d./-]+$", nome):
                    return nome[:120]

        desc = lancamento.descricao or ""
        if "DES:" in desc.upper():
            match = re.search(r"DES:\s*([^\d]{3,})", desc, re.IGNORECASE)
            if match:
                nome = match.group(1).strip()
                if len(nome) > 3:
                    return nome[:120]

        nome = comprovante.arquivo.stem.replace("PAGTO_", "").replace("_", " ").strip()
        if len(nome) > 3:
            return nome[:120]

        return "NÃO IDENTIFICADO"

    def _gerar_referencia(
        self,
        lancamento: LancamentoExtrato,
        comprovante: DocumentoFinanceiro,
    ) -> str:
        refs = [f"L{abs(lancamento.ordem):03d}"]
        if comprovante.identificadores:
            refs.append(comprovante.identificadores[0])
        refs.append(lancamento.data.strftime("%m%Y"))
        return "-".join(refs)


def _criar_documento_financeiro_recibo(
    recibo_path: Path,
    resultado: ResultadoConcilacao,
    generator: ReciboGenerator,
) -> DocumentoFinanceiro:
    """Monta o DocumentoFinanceiro do recibo gerado (sem parse de PDF/DOCX)."""
    favorecido = generator._extrair_favorecido(resultado.comprovante, resultado.lancamento)
    descricao = generator._extrair_descricao(
        resultado.comprovante, resultado.lancamento, favorecido
    )
    return DocumentoFinanceiro(
        arquivo=recibo_path,
        tipo=TipoDocumento.RECIBO,
        valor=resultado.lancamento.valor_absoluto,
        data=resultado.lancamento.data,
        identificadores=[f"L{resultado.lancamento.ordem:03d}"],
        favorecido=favorecido,
        descricao=descricao,
        texto_extraido="",
    )


def gerar_recibos_para_faltantes(
    resultados: List[ResultadoConcilacao],
    pasta_destino: Path,
    modelo_path: Optional[Path] = None,
) -> List[Path]:
    """Gera recibos para PRECISA_RECIBO, vincula documento_base e marca CONCILIADO."""
    generator = ReciboGenerator(modelo_path)
    recibos_gerados: List[Path] = []
    pasta_destino = Path(pasta_destino)
    pasta_destino.mkdir(parents=True, exist_ok=True)

    for resultado in resultados:
        if resultado.status != StatusConciliacao.PRECISA_RECIBO:
            continue
        if not resultado.comprovante:
            continue

        nome = f"RECIBO_{resultado.lancamento.ordem:03d}_{resultado.lancamento.data.strftime('%Y%m%d')}"
        caminho = pasta_destino / nome

        try:
            recibo_path = generator.gerar_recibo(
                resultado.lancamento,
                resultado.comprovante,
                caminho,
            )
            recibos_gerados.append(recibo_path)
            resultado.documento_base = _criar_documento_financeiro_recibo(
                recibo_path, resultado, generator
            )
            resultado.status = StatusConciliacao.CONCILIADO
            resultado.faltantes = []
            resultado.score_base = 100
        except Exception as e:
            print(f"Erro ao gerar recibo para lançamento {resultado.lancamento.ordem}: {e}")

    return recibos_gerados
